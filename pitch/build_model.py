from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference

BLUE = Font(name='Arial', color='0000FF')
BLACK = Font(name='Arial', color='000000')
GREEN = Font(name='Arial', color='008000')
BOLD = Font(name='Arial', bold=True)
H1 = Font(name='Arial', bold=True, size=14, color='0A5C46')
H2 = Font(name='Arial', bold=True, size=11, color='FFFFFF')
SECT = PatternFill('solid', start_color='0A5C46')
YELLOW = PatternFill('solid', start_color='FFFF00')
THIN = Border(bottom=Side(style='thin', color='CCCCCC'))
RS = '"Rs "#,##0;("Rs "#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
USD = '"$"#,##0;("$"#,##0);"-"'

wb = Workbook()

# ── Assumptions ──────────────────────────────────────────────
A = wb.active; A.title = 'Assumptions'
A['A1'] = 'SAHULATKAAR — 5-YEAR MODEL ASSUMPTIONS'; A['A1'].font = H1
A['A2'] = 'Blue cells are inputs — change any of them and the entire model recalculates.'
A['A2'].font = Font(name='Arial', italic=True, size=9, color='666666')

def sect(cell, label):
    A[cell] = label; A[cell].font = H2; A[cell].fill = SECT

def inp(cell, label, value, fmt=NUM, note=None):
    r = int(cell[1:]); A[f'A{r}'] = label; A[f'A{r}'].font = BLACK
    A[cell] = value; A[cell].font = BLUE; A[cell].number_format = fmt
    if note: A[f'H{r}'] = note; A[f'H{r}'].font = Font(name='Arial', size=8, color='888888')

def inp_row(row, label, values, fmt=NUM, note=None):
    A[f'A{row}'] = label; A[f'A{row}'].font = BLACK
    for i, v in enumerate(values):
        c = A.cell(row=row, column=3 + i, value=v); c.font = BLUE; c.number_format = fmt
    if note: A[f'H{row}'] = note; A[f'H{row}'].font = Font(name='Arial', size=8, color='888888')

sect('A4', 'PRICING (Rs / month)')
inp('B5', 'Chhota plan', 2500, RS)
inp('B6', 'Dukaan plan', 5000, RS)
inp('B7', 'Karobaar plan', 12000, RS)
inp('B8', 'Plan mix — Chhota', 0.30, PCT)
inp('B9', 'Plan mix — Dukaan', 0.55, PCT)
inp('B10', 'Plan mix — Karobaar', 0.15, PCT)
A['A11'] = 'Blended ARPU (Rs/mo)'; A['B11'] = '=SUMPRODUCT(B5:B7,B8:B10)'
A['B11'].font = BLACK; A['B11'].number_format = RS
inp('B12', 'Annual price increase (from Y2)', 0.08, PCT)
inp('B13', 'Setup fee per new merchant (one-time)', 5000, RS)

sect('A15', 'GROWTH')
for i, y in enumerate(['Y1', 'Y2', 'Y3', 'Y4', 'Y5']):
    c = A.cell(row=15, column=3 + i, value=y); c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = SECT
inp_row(16, 'New merchants signed / month', [15, 45, 110, 220, 400], NUM, 'Field sales + digital; PK has ~500k WhatsApp-selling retailers')
inp('B17', 'Monthly churn', 0.03, PCT)

sect('A19', 'DIRECT COST PER ACTIVE MERCHANT (Rs / month)')
inp('B20', 'AI (Gemini) usage', 350, RS)
inp('B21', 'Hosting & infrastructure', 120, RS)
inp('B22', 'WhatsApp / BSP fees', 800, RS, 'Blended: BSP early, Meta Tech Provider at scale; customer-initiated msgs are free')
A['A23'] = 'Total direct cost'; A['B23'] = '=SUM(B20:B22)'; A['B23'].font = BLACK; A['B23'].number_format = RS

sect('A25', 'TEAM')
for i, y in enumerate(['Y1', 'Y2', 'Y3', 'Y4', 'Y5']):
    c = A.cell(row=25, column=3 + i, value=y); c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = SECT
inp_row(26, 'Engineers (headcount)', [2, 3, 5, 8, 12])
inp('B27', 'Avg engineer salary (Rs/mo, Y1)', 350000, RS)
inp('B28', 'Merchants per support agent', 150)
inp('B29', 'Support agent salary (Rs/mo, Y1)', 70000, RS)
inp_row(30, 'Management cost (Rs/mo)', [300000, 500000, 800000, 1200000, 1800000], RS)
inp('B31', 'Salary inflation / yr', 0.10, PCT)

sect('A33', 'SALES & MARKETING')
inp('B34', 'CAC per merchant (Rs)', 8000, RS)
inp_row(35, 'Brand budget (Rs/yr)', [1200000, 3000000, 6000000, 10000000, 15000000], RS)
A['A35'] = 'Brand budget (Rs/yr)'

sect('A37', 'OTHER')
inp_row(38, 'G&A / office (Rs/mo)', [150000, 250000, 400000, 700000, 1000000], RS)
inp('B39', 'Corporate tax rate', 0.29, PCT)
inp('B40', 'PKR per USD', 280, NUM)
A['B17'].fill = YELLOW  # churn is the most sensitive assumption
A['B34'].fill = YELLOW

A.column_dimensions['A'].width = 36
for col in 'BCDEFG': A.column_dimensions[col].width = 13
A.column_dimensions['H'].width = 52

# ── Monthly engine (60 months) ───────────────────────────────
M = wb.create_sheet('Monthly')
M['A1'] = 'MONTHLY ENGINE — 60 months (all formulas; drives the annual Forecast)'; M['A1'].font = H1
rows = {3: 'Month #', 4: 'Year', 5: 'Price escalator', 6: 'Blended ARPU (Rs/mo)', 7: 'Merchants — start',
        8: 'New merchants', 9: 'Churned', 10: 'Merchants — end', 11: 'Average active',
        12: 'Subscription revenue (Rs)', 13: 'Setup-fee revenue (Rs)', 14: 'Total revenue (Rs)',
        15: 'Direct costs (Rs)', 16: 'Gross profit (Rs)'}
for r, label in rows.items():
    M[f'A{r}'] = label; M[f'A{r}'].font = BOLD if r in (14, 16) else BLACK
for m in range(1, 61):
    col = get_column_letter(1 + m)
    p = get_column_letter(m)  # previous column
    f = {
        3: m if m == 1 else f'={p}3+1',
        4: f'=INT(({col}3-1)/12)+1',
        5: f'=(1+Assumptions!$B$12)^({col}4-1)',
        6: f'=Assumptions!$B$11*{col}5',
        7: 0 if m == 1 else f'={p}10',
        8: f'=INDEX(Assumptions!$C$16:$G$16,1,{col}4)',
        9: f'={col}7*Assumptions!$B$17',
        10: f'={col}7+{col}8-{col}9',
        11: f'=({col}7+{col}10)/2',
        12: f'={col}11*{col}6',
        13: f'={col}8*Assumptions!$B$13',
        14: f'={col}12+{col}13',
        15: f'={col}11*Assumptions!$B$23',
        16: f'={col}14-{col}15',
    }
    for r, v in f.items():
        c = M[f'{col}{r}']; c.value = v; c.font = BLACK
        c.number_format = NUM if r < 12 else RS
        if r in (5,): c.number_format = '0.000'
M.column_dimensions['A'].width = 26
for m in range(1, 61): M.column_dimensions[get_column_letter(1 + m)].width = 11
M.freeze_panes = 'B3'

# ── Annual Forecast ──────────────────────────────────────────
F = wb.create_sheet('Forecast')
F['A1'] = 'SAHULATKAAR — 5-YEAR FORECAST (Rs)'; F['A1'].font = H1
F['A2'] = 'Every figure is a formula driven by the Assumptions sheet. Simplified EBITDA model (no D&A/financing).'
F['A2'].font = Font(name='Arial', italic=True, size=9, color='666666')
hdr = ['', 'Year 1 (FY27)', 'Year 2 (FY28)', 'Year 3 (FY29)', 'Year 4 (FY30)', 'Year 5 (FY31)']
for i, h in enumerate(hdr):
    c = F.cell(row=3, column=1 + i, value=h); c.font = H2; c.fill = SECT
for i in range(5):
    c = F.cell(row=4, column=2 + i, value=i + 1); c.font = BLACK; c.number_format = '0'
F['A4'] = 'Year index'

def frow(r, label, formula, fmt=RS, bold=False, green=False):
    F[f'A{r}'] = label; F[f'A{r}'].font = BOLD if bold else BLACK
    for i in range(5):
        col = get_column_letter(2 + i)
        c = F[f'{col}{r}']; c.value = formula(col); c.number_format = fmt
        c.font = Font(name='Arial', bold=bold, color='008000' if green else '000000')
        c.border = THIN

frow(6, 'New merchants signed', lambda c: f'=SUMIF(Monthly!$B$4:$BI$4,{c}4,Monthly!$B$8:$BI$8)', NUM, green=True)
frow(7, 'Merchants — end of year', lambda c: f'=INDEX(Monthly!$B$10:$BI$10,1,{c}4*12)', NUM, bold=True, green=True)
frow(8, 'Average active merchants', lambda c: f'=AVERAGEIF(Monthly!$B$4:$BI$4,{c}4,Monthly!$B$11:$BI$11)', NUM, green=True)
frow(10, 'Subscription revenue', lambda c: f'=SUMIF(Monthly!$B$4:$BI$4,{c}4,Monthly!$B$12:$BI$12)', RS, green=True)
frow(11, 'Setup-fee revenue', lambda c: f'=SUMIF(Monthly!$B$4:$BI$4,{c}4,Monthly!$B$13:$BI$13)', RS, green=True)
frow(12, 'Total revenue', lambda c: f'={c}10+{c}11', RS, bold=True)
frow(13, 'Direct costs (AI + hosting + WhatsApp)', lambda c: f'=SUMIF(Monthly!$B$4:$BI$4,{c}4,Monthly!$B$15:$BI$15)', RS, green=True)
frow(14, 'Gross profit', lambda c: f'={c}12-{c}13', RS, bold=True)
frow(15, 'Gross margin', lambda c: f'={c}14/{c}12', PCT)
F['A17'] = 'OPERATING EXPENSES'; F['A17'].font = H2; F['A17'].fill = SECT
frow(18, 'Engineering', lambda c: f'=INDEX(Assumptions!$C$26:$G$26,1,{c}4)*Assumptions!$B$27*12*(1+Assumptions!$B$31)^({c}4-1)')
frow(19, 'Customer support', lambda c: f'=ROUNDUP({c}7/Assumptions!$B$28,0)*Assumptions!$B$29*12*(1+Assumptions!$B$31)^({c}4-1)')
frow(20, 'Management', lambda c: f'=INDEX(Assumptions!$C$30:$G$30,1,{c}4)*12')
frow(21, 'Sales & marketing (CAC × new + brand)', lambda c: f'={c}6*Assumptions!$B$34+INDEX(Assumptions!$C$35:$G$35,1,{c}4)')
frow(22, 'G&A / office', lambda c: f'=INDEX(Assumptions!$C$38:$G$38,1,{c}4)*12')
frow(23, 'Total opex', lambda c: f'=SUM({c}18:{c}22)', RS, bold=True)
frow(25, 'EBITDA', lambda c: f'={c}14-{c}23', RS, bold=True)
frow(26, 'EBITDA margin', lambda c: f'=IF({c}12=0,0,{c}25/{c}12)', PCT)
frow(27, 'Tax (on positive EBITDA)', lambda c: f'=MAX(0,{c}25)*Assumptions!$B$39')
frow(28, 'Net result', lambda c: f'={c}25-{c}27', RS, bold=True)
F['A29'] = 'Cumulative net cash'; F['A29'].font = BOLD
F['B29'] = '=B28'; 
for i in range(1, 5):
    col = get_column_letter(2 + i); prev = get_column_letter(1 + i)
    F[f'{col}29'] = f'={prev}29+{col}28'
for i in range(5):
    c = F[f'{get_column_letter(2+i)}29']; c.number_format = RS; c.font = BOLD
F['A31'] = 'IN USD'; F['A31'].font = H2; F['A31'].fill = SECT
frow(32, 'Revenue (USD)', lambda c: f'={c}12/Assumptions!$B$40', USD)
frow(33, 'EBITDA (USD)', lambda c: f'={c}25/Assumptions!$B$40', USD)
frow(35, 'ARR run-rate at year end', lambda c: f'=INDEX(Monthly!$B$10:$BI$10,1,{c}4*12)*INDEX(Monthly!$B$6:$BI$6,1,{c}4*12)*12', RS)
F['A37'] = 'UNIT ECONOMICS'; F['A37'].font = H2; F['A37'].fill = SECT
F['A38'] = 'LTV (ARPU × GM ÷ churn)'; F['B38'] = '=Assumptions!B11*Forecast!B15/Assumptions!B17'
F['B38'].number_format = RS; F['B38'].font = BLACK
F['A39'] = 'LTV / CAC'; F['B39'] = '=B38/Assumptions!B34'; F['B39'].number_format = '0.0"x"'; F['B39'].font = BLACK
F['A40'] = 'CAC payback (months)'; F['B40'] = '=Assumptions!B34/(Assumptions!B11*Forecast!B15)'
F['B40'].number_format = '0.0'; F['B40'].font = BLACK
F.column_dimensions['A'].width = 38
for col in 'BCDEF': F.column_dimensions[col].width = 16

# ── Dashboard with charts ────────────────────────────────────
D = wb.create_sheet('Dashboard')
D['A1'] = 'SAHULATKAAR — INVESTOR DASHBOARD'; D['A1'].font = H1
kpis = [('Merchants (Year 5)', '=Forecast!F7', NUM), ('Revenue Year 5', '=Forecast!F12', RS),
        ('Revenue Year 5 (USD)', '=Forecast!F32', USD), ('EBITDA margin Year 5', '=Forecast!F26', PCT),
        ('Gross margin', '=Forecast!F15', PCT), ('LTV / CAC', '=Forecast!B39', '0.0"x"'),
        ('CAC payback (months)', '=Forecast!B40', '0.0')]
for i, (label, f, fmt) in enumerate(kpis):
    D.cell(row=3, column=1 + i, value=label).font = BOLD
    c = D.cell(row=4, column=1 + i, value=f); c.font = GREEN; c.number_format = fmt
    D.column_dimensions[get_column_letter(1 + i)].width = 20

ch1 = BarChart(); ch1.type = 'col'; ch1.title = 'Revenue vs EBITDA (Rs)'; ch1.height = 9; ch1.width = 22
data = Reference(F, min_col=1, min_row=12, max_col=6, max_row=12)
ch1.add_data(data, titles_from_data=True, from_rows=True)
data2 = Reference(F, min_col=1, min_row=25, max_col=6, max_row=25)
ch1.add_data(data2, titles_from_data=True, from_rows=True)
cats = Reference(F, min_col=2, min_row=3, max_col=6, max_row=3)
ch1.set_categories(cats)
D.add_chart(ch1, 'A7')

ch2 = LineChart(); ch2.title = 'Active merchants (end of year)'; ch2.height = 9; ch2.width = 22
d3 = Reference(F, min_col=1, min_row=7, max_col=6, max_row=7)
ch2.add_data(d3, titles_from_data=True, from_rows=True)
ch2.set_categories(cats)
D.add_chart(ch2, 'A27')

wb.save('Sahulatkaar-Financial-Model-5yr.xlsx')
print('saved')
